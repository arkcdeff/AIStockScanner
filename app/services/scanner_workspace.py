from __future__ import annotations

import csv
import io
from datetime import date
from typing import Any

from fpdf import FPDF
from openpyxl import Workbook
from sqlalchemy import ColumnElement, Select, asc, desc, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.daily_price import DailyPrice
from app.models.indicator_result import IndicatorResult
from app.models.scan_definition import ScanDefinition
from app.models.scan_result import ScanResult
from app.models.sector import Sector
from app.models.stock import Stock


class ScannerWorkspaceService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def get_filter_options(self) -> dict[str, Any]:
        sectors = (
            await self.session.scalars(select(Sector.name).order_by(Sector.name))
        ).all()
        scanners = (
            await self.session.execute(
                select(ScanDefinition.code, ScanDefinition.name).where(
                    ScanDefinition.is_active.is_(True)
                ).order_by(ScanDefinition.name)
            )
        ).fetchall()
        scanner_map = {code: name for code, name in scanners}

        dates = (
            await self.session.scalars(
                select(ScanResult.scan_date).distinct().order_by(ScanResult.scan_date.desc()).limit(30)
            )
        ).all()

        return {
            "sectors": list(sectors),
            "scanners": scanner_map,
            "dates": [d.isoformat() if isinstance(d, date) else str(d) for d in dates],
        }

    async def get_results(
        self,
        *,
        search: str = "",
        scanner: str | None = None,
        sector: str | None = None,
        date_from: str | None = None,
        date_to: str | None = None,
        min_price: float | None = None,
        max_price: float | None = None,
        min_delivery: float | None = None,
        max_delivery: float | None = None,
        min_volume: int | None = None,
        max_volume: int | None = None,
        sort_by: str = "scan_date",
        sort_dir: str = "desc",
        page: int = 1,
        page_size: int = 25,
    ) -> dict[str, Any]:
        statement = self._build_query(
            search=search,
            scanner=scanner,
            sector=sector,
            date_from=date_from,
            date_to=date_to,
            min_price=min_price,
            max_price=max_price,
            min_delivery=min_delivery,
            max_delivery=max_delivery,
            min_volume=min_volume,
            max_volume=max_volume,
            sort_by=sort_by,
            sort_dir=sort_dir,
        )

        count_stmt = select(func.count()).select_from(statement.subquery())
        total = (await self.session.scalar(count_stmt)) or 0

        offset = (page - 1) * page_size
        paginated = select("*").select_from(statement.subquery()).offset(offset).limit(page_size)
        rows = (await self.session.execute(paginated)).fetchall()

        results = [self._row_to_dict(row) for row in rows]
        return {
            "results": results,
            "total": total,
            "page": page,
            "page_size": page_size,
            "pages": (total + page_size - 1) // page_size if page_size else 0,
        }

    async def get_export_data(
        self,
        *,
        search: str = "",
        scanner: str | None = None,
        sector: str | None = None,
        date_from: str | None = None,
        date_to: str | None = None,
        min_price: float | None = None,
        max_price: float | None = None,
        min_delivery: float | None = None,
        max_delivery: float | None = None,
        min_volume: int | None = None,
        max_volume: int | None = None,
        sort_by: str = "scan_date",
        sort_dir: str = "desc",
    ) -> list[dict[str, Any]]:
        statement = self._build_query(
            search=search,
            scanner=scanner,
            sector=sector,
            date_from=date_from,
            date_to=date_to,
            min_price=min_price,
            max_price=max_price,
            min_delivery=min_delivery,
            max_delivery=max_delivery,
            min_volume=min_volume,
            max_volume=max_volume,
            sort_by=sort_by,
            sort_dir=sort_dir,
        )
        rows = (await self.session.execute(statement)).fetchall()
        return [self._row_to_dict(row) for row in rows]

    def _build_query(
        self,
        *,
        search: str,
        scanner: str | None,
        sector: str | None,
        date_from: str | None,
        date_to: str | None,
        min_price: float | None,
        max_price: float | None,
        min_delivery: float | None,
        max_delivery: float | None,
        min_volume: int | None,
        max_volume: int | None,
        sort_by: str,
        sort_dir: str,
    ) -> Select:
        price_history = (
            select(
                DailyPrice.stock_id,
                DailyPrice.trade_date,
                DailyPrice.volume,
                func.avg(DailyPrice.volume).over(
                    partition_by=DailyPrice.stock_id,
                    order_by=DailyPrice.trade_date.asc(),
                    rows=(-20, -1),
                ).label("avg_volume_20d"),
            )
            .subquery()
        )

        stmt = (
            select(
                ScanResult.scan_date,
                ScanResult.score.label("scanner_score"),
                ScanResult.metrics.label("metrics"),
                ScanDefinition.code.label("scanner_code"),
                ScanDefinition.name.label("scanner_name"),
                ScanDefinition.category.label("scanner_category"),
                Stock.symbol,
                Stock.company_name,
                Sector.name.label("sector"),
                DailyPrice.close.label("price"),
                DailyPrice.volume,
                DailyPrice.delivery_percentage,
                IndicatorResult.rsi_14,
                IndicatorResult.macd,
                IndicatorResult.adx,
                IndicatorResult.ema_20,
                IndicatorResult.ema_50,
                IndicatorResult.ema_200,
                price_history.c.avg_volume_20d,
            )
            .join(ScanDefinition, ScanResult.scan_definition_id == ScanDefinition.id)
            .join(Stock, ScanResult.stock_id == Stock.id)
            .join(Sector, Stock.sector_id == Sector.id, isouter=True)
            .join(DailyPrice, (DailyPrice.stock_id == Stock.id) & (DailyPrice.trade_date == ScanResult.scan_date))
            .join(IndicatorResult, (IndicatorResult.stock_id == Stock.id) & (IndicatorResult.trade_date == ScanResult.scan_date))
            .join(price_history, (price_history.c.stock_id == Stock.id) & (price_history.c.trade_date == ScanResult.scan_date), isouter=True)
            .where(True)
        )

        if search:
            stmt = stmt.where(
                or_(
                    Stock.symbol.ilike(f"%{search}%"),
                    Stock.company_name.ilike(f"%{search}%"),
                    ScanDefinition.name.ilike(f"%{search}%"),
                )
            )

        if scanner:
            stmt = stmt.where(ScanDefinition.code == scanner)

        if sector:
            stmt = stmt.where(Sector.name == sector)

        if date_from:
            stmt = stmt.where(ScanResult.scan_date >= date_from)

        if date_to:
            stmt = stmt.where(ScanResult.scan_date <= date_to)

        if min_price is not None:
            stmt = stmt.where(DailyPrice.close >= min_price)

        if max_price is not None:
            stmt = stmt.where(DailyPrice.close <= max_price)

        if min_delivery is not None:
            stmt = stmt.where(DailyPrice.delivery_percentage >= min_delivery)

        if max_delivery is not None:
            stmt = stmt.where(DailyPrice.delivery_percentage <= max_delivery)

        if min_volume is not None:
            stmt = stmt.where(DailyPrice.volume >= min_volume)

        if max_volume is not None:
            stmt = stmt.where(DailyPrice.volume <= max_volume)

        sort_column = {
            "symbol": Stock.symbol,
            "company": Stock.company_name,
            "sector": Sector.name,
            "price": DailyPrice.close,
            "volume": DailyPrice.volume,
            "delivery": DailyPrice.delivery_percentage,
            "rsi": IndicatorResult.rsi_14,
            "macd": IndicatorResult.macd,
            "adx": IndicatorResult.adx,
            "scanner_score": ScanResult.score,
            "date": ScanResult.scan_date,
            "scanner": ScanDefinition.name,
        }.get(sort_by, ScanResult.scan_date)

        if sort_dir == "asc":
            stmt = stmt.order_by(asc(sort_column))
        else:
            stmt = stmt.order_by(desc(sort_column))

        return stmt

    def _row_to_dict(self, row: Any) -> dict[str, Any]:
        avg_volume_20d = getattr(row, "avg_volume_20d", None)
        volume = getattr(row, "volume", None)
        volume_multiple = None
        if avg_volume_20d and volume and avg_volume_20d > 0:
            volume_multiple = round(float(volume) / float(avg_volume_20d), 2)

        ema_trend = "Neutral"
        ema_20 = getattr(row, "ema_20", None)
        ema_50 = getattr(row, "ema_50", None)
        ema_200 = getattr(row, "ema_200", None)
        if ema_20 and ema_50 and ema_200:
            if float(ema_20) > float(ema_50) > float(ema_200):
                ema_trend = "Bullish"
            elif float(ema_20) < float(ema_50) < float(ema_200):
                ema_trend = "Bearish"

        metrics = getattr(row, "metrics", {}) or {}
        reason = metrics.get("reason", "") if isinstance(metrics, dict) else ""

        scan_date = getattr(row, "scan_date", None)
        if hasattr(scan_date, "isoformat"):
            scan_date = scan_date.isoformat()

        return {
            "symbol": getattr(row, "symbol", ""),
            "company_name": getattr(row, "company_name", ""),
            "sector": getattr(row, "sector", "") or "N/A",
            "price": float(getattr(row, "price", 0) or 0),
            "volume": int(getattr(row, "volume", 0) or 0),
            "avg_volume_20d": float(avg_volume_20d) if avg_volume_20d is not None else None,
            "volume_multiple": volume_multiple,
            "delivery": float(getattr(row, "delivery_percentage", 0) or 0),
            "rsi": float(getattr(row, "rsi_14", 0) or 0),
            "macd": float(getattr(row, "macd", 0) or 0),
            "adx": float(getattr(row, "adx", 0) or 0),
            "ema_trend": ema_trend,
            "scanner_score": float(getattr(row, "scanner_score", 0) or 0),
            "reason": reason,
            "scan_date": scan_date,
            "scanner_name": getattr(row, "scanner_name", ""),
            "scanner_code": getattr(row, "scanner_code", ""),
        }

    def to_csv(self, rows: list[dict[str, Any]]) -> bytes:
        output = io.StringIO()
        writer = csv.DictWriter(
            output,
            fieldnames=[
                "symbol", "company_name", "sector", "price", "volume",
                "avg_volume_20d", "volume_multiple", "delivery", "rsi",
                "macd", "adx", "ema_trend", "scanner_score", "reason", "scan_date",
            ],
        )
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in writer.fieldnames})
        return output.getvalue().encode("utf-8")

    def to_excel(self, rows: list[dict[str, Any]]) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Scanner Results"
        headers = [
            "Symbol", "Company", "Sector", "Price", "Volume",
            "20D Avg Volume", "Volume Multiple", "Delivery", "RSI",
            "MACD", "ADX", "EMA Trend", "Scanner Score", "Reason", "Scan Date",
        ]
        ws.append(headers)
        for row in rows:
            ws.append([
                row.get("symbol", ""),
                row.get("company_name", ""),
                row.get("sector", ""),
                row.get("price", 0),
                row.get("volume", 0),
                row.get("avg_volume_20d", ""),
                row.get("volume_multiple", ""),
                row.get("delivery", 0),
                row.get("rsi", 0),
                row.get("macd", 0),
                row.get("adx", 0),
                row.get("ema_trend", ""),
                row.get("scanner_score", 0),
                row.get("reason", ""),
                row.get("scan_date", ""),
            ])
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def to_pdf(self, rows: list[dict[str, Any]]) -> bytes:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=10)
        headers = [
            "Symbol", "Company", "Sector", "Price", "Volume",
            "20D Avg", "Vol Mult", "Delivery", "RSI", "MACD", "ADX",
            "EMA", "Score", "Reason", "Date",
        ]
        col_widths = [18, 28, 18, 16, 18, 18, 16, 16, 12, 16, 12, 12, 12, 35, 18]

        pdf.set_font("Helvetica", style="B", size=9)
        for i, header in enumerate(headers):
            pdf.cell(col_widths[i], 6, header, border=1, align="C")
        pdf.ln()

        pdf.set_font("Helvetica", size=8)
        for row in rows:
            values = [
                str(row.get("symbol", ""))[:18],
                str(row.get("company_name", ""))[:28],
                str(row.get("sector", ""))[:18],
                f"{row.get('price', 0):.2f}",
                str(row.get("volume", 0)),
                f"{row.get('avg_volume_20d', 0):.0f}" if row.get("avg_volume_20d") else "",
                str(row.get("volume_multiple", "")) if row.get("volume_multiple") else "",
                f"{row.get('delivery', 0):.1f}%",
                f"{row.get('rsi', 0):.1f}",
                f"{row.get('macd', 0):.2f}",
                f"{row.get('adx', 0):.1f}",
                str(row.get("ema_trend", ""))[:12],
                f"{row.get('scanner_score', 0):.2f}",
                str(row.get("reason", ""))[:35],
                str(row.get("scan_date", ""))[:18],
            ]
            for i, value in enumerate(values):
                pdf.cell(col_widths[i], 6, value, border=1, align="L")
            pdf.ln()

        return bytes(pdf.output(dest="S"))
